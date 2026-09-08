import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List

from docx import Document
from docx.shared import Pt, Inches, RGBColor

from sovereign_ai.core.config import DELIVERABLES_DIR, TASKS_DIR
from sovereign_ai.core.logger import logger, SovereignException
from sovereign_ai.core.provenance import Provenance
from sovereign_ai.schemas.security import UserContext, UserRole


class ControlledDeliverableService:
    """
    Controlled deliverable service for sovereign on-premise AI inspection reports.
    Enforces deterministic human gating, local DOCX rendering, SHA-256 verification,
    and immutable provenance tracking.
    """

    def __init__(self):
        self.deliverables_dir = DELIVERABLES_DIR
        self.tasks_dir = TASKS_DIR
        self.provenance = Provenance()
        self.deliverables_dir.mkdir(parents=True, exist_ok=True)
        self.tasks_dir.mkdir(parents=True, exist_ok=True)

    # -------------------------------------------------------------------------
    # Task State Persistence
    # -------------------------------------------------------------------------
    def save_task_state(self, task_id: str, state: Dict[str, Any]) -> None:
        """
        Persists task execution state atomically to TASKS_DIR/{task_id}.json.
        Survives process restarts.
        """
        if not task_id:
            raise SovereignException("Task ID is required for persistence.", status_code=400)

        filepath = self.tasks_dir / f"{task_id}.json"
        state_to_save = dict(state)
        state_to_save["task_id"] = task_id
        state_to_save["updated_at"] = datetime.now().isoformat()
        if "created_at" not in state_to_save:
            state_to_save["created_at"] = datetime.now().isoformat()

        # Write to temporary file first and atomically rename
        tmp_path = self.tasks_dir / f"{task_id}.tmp"
        try:
            with open(tmp_path, "w", encoding="utf-8") as f:
                json.dump(state_to_save, f, indent=2, default=str)
            tmp_path.replace(filepath)
        except Exception as e:
            if tmp_path.exists():
                tmp_path.unlink()
            logger.error(f"Failed to persist task state for {task_id}: {e}")
            raise SovereignException(f"Failed to persist task state: {str(e)}", status_code=500)

    def get_task_state(self, task_id: str) -> Optional[Dict[str, Any]]:
        """
        Loads persisted task execution state from disk.
        """
        filepath = self.tasks_dir / f"{task_id}.json"
        if not filepath.exists():
            return None
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logger.error(f"Failed to load task state for {task_id}: {e}")
            raise SovereignException(f"Corrupt task state for task '{task_id}': {str(e)}", status_code=500)

    def update_task_state(self, task_id: str, updates: Dict[str, Any]) -> Dict[str, Any]:
        """
        Merges updates into existing persisted task state.
        """
        state = self.get_task_state(task_id)
        if not state:
            raise SovereignException(f"Task '{task_id}' not found.", status_code=404)
        state.update(updates)
        self.save_task_state(task_id, state)
        return state

    # -------------------------------------------------------------------------
    # Deliverable Generation & Gating Enforcement
    # -------------------------------------------------------------------------
    def generate_deliverable(
        self,
        task_record: Dict[str, Any],
        caller_user: Optional[UserContext] = None,
    ) -> Dict[str, Any]:
        """
        Validates risk and human gate status, then generates a controlled DOCX deliverable.
        Strictly refuses delivery if:
        - risk is MEDIUM or HIGH AND human_gate is PENDING
        - human_gate is REJECTED
        - human_gate is EDIT_REQUIRED
        """
        task_id = task_record.get("task_id")
        if not task_id:
            raise SovereignException("Task record missing required 'task_id'.", status_code=400)

        risk = task_record.get("risk", {})
        risk_level = risk.get("risk_level", "HIGH")
        human_gate = task_record.get("human_gate", {})
        gate_status = human_gate.get("human_gate_status", "PENDING")
        human_approval = human_gate.get("human_approval", False)

        # 1. Check Gating Policies
        if gate_status == "REJECTED":
            self.provenance.log_deliverable_event(
                event_type="DELIVERY_BLOCKED",
                task_id=task_id,
                status="BLOCKED",
                actor=caller_user.username if caller_user else None,
                role=caller_user.role.value if caller_user else None,
                reason="Task was REJECTED by reviewer. Controlled deliverable generation is prohibited.",
            )
            raise SovereignException(
                f"Delivery blocked: Task '{task_id}' was REJECTED. No deliverable may be generated.",
                status_code=403,
            )

        if gate_status == "EDIT_REQUIRED":
            self.provenance.log_deliverable_event(
                event_type="DELIVERY_BLOCKED",
                task_id=task_id,
                status="BLOCKED",
                actor=caller_user.username if caller_user else None,
                role=caller_user.role.value if caller_user else None,
                reason="Task requires edits before approval. Controlled deliverable generation is held.",
            )
            raise SovereignException(
                f"Delivery blocked: Task '{task_id}' is in EDIT_REQUIRED state. Changes must be re-evaluated and approved before deliverable generation.",
                status_code=400,
            )

        if risk_level in ["MEDIUM", "HIGH"] and gate_status != "APPROVED":
            self.provenance.log_deliverable_event(
                event_type="DELIVERY_BLOCKED",
                task_id=task_id,
                status="BLOCKED",
                actor=caller_user.username if caller_user else None,
                role=caller_user.role.value if caller_user else None,
                reason=f"Task has {risk_level} risk and gate status '{gate_status}'. Explicit human approval is required.",
            )
            raise SovereignException(
                f"Delivery blocked: Task '{task_id}' has risk '{risk_level}' with human gate status '{gate_status}'. Controlled deliverable requires explicit human approval.",
                status_code=403,
            )

        # 2. Authorization Confirmed -> Log DELIVERY_REQUESTED
        self.provenance.log_deliverable_event(
            event_type="DELIVERY_REQUESTED",
            task_id=task_id,
            status="IN_PROGRESS",
            actor=caller_user.username if caller_user else None,
            role=caller_user.role.value if caller_user else None,
            details={"risk_level": risk_level, "gate_status": gate_status},
        )

        # 3. Construct DOCX Report
        docx_filename = f"inspection_report_{task_id}.docx"
        docx_path = self.deliverables_dir / docx_filename

        try:
            self._render_docx(docx_path=docx_path, task_record=task_record)
        except Exception as e:
            logger.error(f"DOCX rendering failed for {task_id}: {e}")
            self.provenance.log_deliverable_event(
                event_type="DELIVERABLE_GENERATED",
                task_id=task_id,
                status="FAILED",
                reason=f"DOCX rendering exception: {str(e)}",
            )
            raise SovereignException(f"Failed to generate deliverable DOCX: {str(e)}", status_code=500)

        # 4. Calculate SHA-256 over Final File Bytes
        if not docx_path.exists() or docx_path.stat().st_size == 0:
            raise SovereignException("DOCX file generation failed or produced an empty file.", status_code=500)

        file_bytes = docx_path.read_bytes()
        file_sha256 = hashlib.sha256(file_bytes).hexdigest()

        # 5. Build Deliverable Metadata
        deliverable_meta = {
            "status": "delivered",
            "path": str(docx_path.resolve()),
            "filename": docx_filename,
            "sha256": file_sha256,
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "generated_at": datetime.now().isoformat(),
            "file_size_bytes": len(file_bytes),
        }

        # 6. Stream DELIVERABLE_GENERATED Event
        self.provenance.log_deliverable_event(
            event_type="DELIVERABLE_GENERATED",
            task_id=task_id,
            status="SUCCESS",
            actor=caller_user.username if caller_user else None,
            role=caller_user.role.value if caller_user else None,
            output_hash=file_sha256,
            details={
                "filename": docx_filename,
                "file_size": len(file_bytes),
                "risk_level": risk_level,
                "gate_status": gate_status,
            },
        )

        return deliverable_meta

    # -------------------------------------------------------------------------
    # Local DOCX Report Layout
    # -------------------------------------------------------------------------
    def _render_docx(self, docx_path: Path, task_record: Dict[str, Any]) -> None:
        """
        Creates a structured, controlled DOCX document using python-docx.
        Strictly includes verified findings, RAG evidence, verification breakdown,
        risk assessment, provenance hashes, and audit lineage.
        Does not invent facts or measurements.
        """
        doc = Document()

        # Page setup
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Helper for adding styled headings
        def add_heading_1(text: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(30, 41, 59)  # Slate 800

        def add_divider():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 50)
            run.font.color.rgb = RGBColor(148, 163, 184)  # Slate 400

        # Title Block
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(2)
        r1 = title_p.add_run("SOVEREIGN AI WORKBENCH\n")
        r1.bold = True
        r1.font.size = Pt(18)
        r1.font.color.rgb = RGBColor(15, 23, 42)

        r2 = title_p.add_run("CONTROLLED INSPECTION ANALYSIS REPORT")
        r2.bold = True
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(71, 85, 105)

        add_divider()

        # Metadata Header Table
        meta_table = doc.add_table(rows=6, cols=2)
        meta_table.style = "Table Grid"

        meta_rows = [
            ("Task ID:", task_record.get("task_id", "N/A")),
            ("Document / Context:", task_record.get("audit_record", {}).get("filename") or task_record.get("filename") or "N/A"),
            ("Security Classification:", task_record.get("security", {}).get("classification", "GENERAL")),
            ("Risk Assessment Level:", task_record.get("risk", {}).get("risk_level", "LOW")),
            ("Human Gate Status:", task_record.get("human_gate", {}).get("human_gate_status", "NOT_REQUIRED")),
            ("Human Approval Authorization:", "APPROVED" if task_record.get("human_gate", {}).get("human_approval") else "NOT_APPROVED (Automated / Pending)"),
        ]

        for i, (label, val) in enumerate(meta_rows):
            cell_label = meta_table.cell(i, 0)
            cell_val = meta_table.cell(i, 1)
            cell_label.paragraphs[0].text = label
            cell_label.paragraphs[0].runs[0].bold = True
            cell_val.paragraphs[0].text = str(val)

        add_divider()

        # 1. EXECUTIVE SUMMARY
        add_heading_1("1. EXECUTIVE SUMMARY")
        answer_text = task_record.get("answer", "No verified findings available.")
        p_exec = doc.add_paragraph(answer_text)
        p_exec.paragraph_format.space_after = Pt(10)

        # 2. EVIDENCE / RETRIEVED KNOWLEDGE
        add_heading_1("2. EVIDENCE / RETRIEVED KNOWLEDGE")
        retrieved = task_record.get("retrieved_knowledge", [])
        if not retrieved and "stages" in task_record:
            # Check for retrieved sources in stages
            for stage in task_record.get("stages", []):
                if stage.get("stage") == "RETRIEVE":
                    retrieved = stage.get("details", {}).get("sources", [])
                    break

        if retrieved:
            for idx, item in enumerate(retrieved, start=1):
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_after = Pt(4)
                chunk_id = item.get("chunk_id", f"Source #{idx}")
                source_name = item.get("source", "Enterprise Knowledge Store")
                classification = item.get("classification", "GENERAL")
                score = item.get("similarity_score", "N/A")

                r = p_item.add_run(f"• Source {idx}: {chunk_id} ({source_name})\n")
                r.bold = True
                p_item.add_run(f"  Classification: {classification} | Similarity Distance: {score}")
        else:
            p_none = doc.add_paragraph("No external organizational knowledge sources were retrieved for this task.")
            p_none.runs[0].italic = True

        # 3. VERIFICATION
        add_heading_1("3. VERIFICATION")
        v = task_record.get("verification", {})
        p_ver = doc.add_paragraph()
        p_ver.paragraph_format.space_after = Pt(6)
        p_ver.add_run("• Verification Status: ").bold = True
        p_ver.add_run(f"{v.get('status', 'UNKNOWN')}\n")
        p_ver.add_run("• Factual Grounding: ").bold = True
        p_ver.add_run(f"{v.get('grounded', False)}\n")
        p_ver.add_run("• Evidence Sources Corroborated: ").bold = True
        p_ver.add_run(f"{v.get('evidence_sources', 0)}\n")
        p_ver.add_run("• Missing Evidence Flag: ").bold = True
        p_ver.add_run(f"{v.get('missing_evidence', False)}")

        unsupported = v.get("unsupported_claims", [])
        if unsupported:
            doc.add_paragraph("Unsupported Claims / Caveats:").runs[0].bold = True
            for claim in unsupported:
                doc.add_paragraph(f"  - {claim}")

        # 4. RISK ASSESSMENT
        add_heading_1("4. RISK ASSESSMENT")
        risk = task_record.get("risk", {})
        p_risk = doc.add_paragraph()
        p_risk.add_run(f"Evaluated Risk Level: {risk.get('risk_level', 'UNKNOWN')}\n").bold = True
        p_risk.add_run(f"Operational Impact: {risk.get('operational_impact', 'LOW')}")

        factors = risk.get("factors", [])
        if factors:
            doc.add_paragraph("Risk Factors Evaluated:").runs[0].bold = True
            for factor in factors:
                doc.add_paragraph(f"  - {factor}")

        # 5. PROVENANCE & GOVERNANCE
        add_heading_1("5. PROVENANCE & CRYPTOGRAPHIC LINEAGE")
        audit_rec = task_record.get("audit_record", {})
        model_info = task_record.get("model", {})
        p_prov = doc.add_paragraph()
        p_prov.paragraph_format.space_after = Pt(6)
        p_prov.add_run("• Input SHA-256: ").bold = True
        p_prov.add_run(f"{audit_rec.get('input_sha256', 'N/A')}\n")
        p_prov.add_run("• Inference Model: ").bold = True
        p_prov.add_run(f"{model_info.get('model', audit_rec.get('model', 'N/A'))}\n")
        p_prov.add_run("• Verification State: ").bold = True
        p_prov.add_run(f"{v.get('status', 'N/A')}\n")
        p_prov.add_run("• Human Gate Authorization: ").bold = True
        p_prov.add_run(f"{task_record.get('human_gate', {}).get('human_gate_status', 'N/A')}")

        # 6. AUDIT
        add_heading_1("6. AUDIT & DELIVERY TRAIL")
        gate_info = task_record.get("human_gate", {})
        p_audit = doc.add_paragraph()
        p_audit.add_run("• Workflow Task Identifier: ").bold = True
        p_audit.add_run(f"{task_record.get('task_id', 'N/A')}\n")
        p_audit.add_run("• Delivery Status: ").bold = True
        p_audit.add_run(f"{task_record.get('delivery', {}).get('delivery_status', 'DELIVERED')}\n")
        p_audit.add_run("• Timestamp Generated: ").bold = True
        p_audit.add_run(f"{datetime.now().isoformat()}\n")

        if gate_info.get("approved_by"):
            p_audit.add_run("• Authorized Approver: ").bold = True
            p_audit.add_run(f"{gate_info.get('approved_by')} at {gate_info.get('approved_at')}")

        doc.save(str(docx_path))
